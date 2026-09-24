from pathlib import Path
import re

import docx
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentProcessor:
    supported_formats = {"pdf", "docx", "txt", "xlsx"}

    def _get_filename(self, file_object, filename=None):
        if filename:
            return Path(filename).name
        if isinstance(file_object, (str, Path)):
            return Path(file_object).name
        uploaded_name = getattr(file_object, "name", None)
        if not uploaded_name:
            raise ValueError("Uploaded document is missing a filename")
        return Path(uploaded_name).name

    def clean_text(self, text):
        return re.sub(r"\s+", " ", text or "").strip()

    def extract_text_from_pdf(self, file_object):
        reader = PdfReader(file_object)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def extract_text_from_docx(self, file_object):
        document = docx.Document(file_object)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)

    def extract_text_from_txt(self, file_object):
        data = file_object.read()
        if isinstance(data, str):
            return data
        return data.decode("utf-8", errors="replace")

    def extract_text_from_xlsx(self, file_object):
        workbook = load_workbook(file_object, read_only=True, data_only=True)
        rows = []
        for worksheet in workbook.worksheets:
            rows.append(f"Sheet: {worksheet.title}")
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
        return "\n".join(rows)

    def process_document(self, file_object, filename=None):
        filename = self._get_filename(file_object, filename)
        extension = Path(filename).suffix.lower().lstrip(".")
        if extension not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {extension or 'unknown'}")

        if isinstance(file_object, (str, Path)):
            with open(file_object, "rb") as handle:
                return self.process_document(handle, filename)

        if hasattr(file_object, "seek"):
            file_object.seek(0)
        extractors = {
            "pdf": self.extract_text_from_pdf,
            "docx": self.extract_text_from_docx,
            "txt": self.extract_text_from_txt,
            "xlsx": self.extract_text_from_xlsx,
        }
        return self.clean_text(extractors[extension](file_object))

    def process_documents(self, files):
        processed = []
        paths = [Path(files)] if isinstance(files, (str, Path)) else list(files)
        if len(paths) == 1 and isinstance(paths[0], Path) and paths[0].is_dir():
            paths = [path for path in paths[0].iterdir() if path.is_file()]

        for file_object in paths:
            filename = self._get_filename(file_object)
            extension = Path(filename).suffix.lower().lstrip(".")
            if extension not in self.supported_formats:
                continue
            try:
                text = self.process_document(file_object, filename)
            except Exception as error:
                raise ValueError(f"Unable to process {filename}: {error}") from error
            if text:
                processed.append({
                    "text": text,
                    "metadata": {"source": filename, "file_type": extension},
                })
        return processed


def process_documents(files):
    """Process Streamlit uploads or a directory into indexable documents."""
    return DocumentProcessor().process_documents(files)